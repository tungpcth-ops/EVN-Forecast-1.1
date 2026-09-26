import io
import pandas as pd
from docx import Document
from docx.shared import Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle


def _fmt(x):
    try:
        return f"{float(x):,.0f}".replace(",", ".")
    except Exception:
        return str(x)


def build_docx(series, forecasts, backtest, top100=None, state=None, outage_monthly=None, outage_details=None, pcth_detail=None, pcth_summary=None, pcth_scenarios=None, pcth_audit=None, title="BÁO CÁO EVN FORECAST 1.3 PCTH"):
    bio = io.BytesIO()
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    p = doc.add_paragraph(); p.alignment = 1
    r = p.add_run(title); r.bold = True; r.font.size = Pt(15)
    p = doc.add_paragraph(); p.alignment = 1; p.add_run("Điện lực Thường Xuân").bold = True

    doc.add_heading("1. Tổng quan", level=1)
    latest = series.iloc[-1]
    doc.add_paragraph(f"Tháng dữ liệu mới nhất: {pd.to_datetime(latest['date']).strftime('%m/%Y')}")
    doc.add_paragraph(f"Điện thương phẩm thực tế: {_fmt(latest['actual'])} kWh")
    if "normalized_actual" in latest.index:
        doc.add_paragraph(f"Điện thương phẩm chuẩn hóa sau ảnh hưởng mất điện: {_fmt(latest['normalized_actual'])} kWh")
        doc.add_paragraph(f"Điện năng ước không thực hiện do mất điện: {_fmt(latest.get('outage_lost_kwh',0))} kWh")
    if state:
        doc.add_paragraph(f"Model State: EVNF_{state.get('latest_month') or 'NA'}_v1.2")
        doc.add_paragraph(f"Cập nhật: {state.get('updated_at') or 'Chưa lưu'}")

    if outage_monthly is not None and not outage_monthly.empty:
        doc.add_heading("2. Phân tích mất điện", level=1)
        t = doc.add_table(rows=1, cols=2)
        t.rows[0].cells[0].text = "Tháng"; t.rows[0].cells[1].text = "Điện năng ước không thực hiện (kWh)"
        for _, row in outage_monthly.iterrows():
            cells=t.add_row().cells
            cells[0].text=pd.to_datetime(row['date']).strftime('%m/%Y')
            cells[1].text=_fmt(row['lost_kwh'])
        if outage_details is not None and not outage_details.empty:
            doc.add_paragraph(f"Số dòng KH-sự kiện mất điện đã phân tích: {len(outage_details):,}")
            doc.add_paragraph(f"Tổng KH-giờ mất điện: {outage_details['duration_hours'].fillna(0).sum():,.2f} giờ (tính theo từng dòng KH-sự kiện).")


    if pcth_summary is not None and not pcth_summary.empty:
        doc.add_heading("3. PCTH Forecast Mode", level=1)
        r0=pcth_summary.iloc[0]
        doc.add_paragraph(f"Tổng nền: {_fmt(r0.get('Tổng nền (kWh)',0))} kWh")
        doc.add_paragraph(f"Tổng dự báo PCTH: {_fmt(r0.get('Tổng dự báo PCTH (kWh)',0))} kWh")
        doc.add_paragraph(f"Tăng/giảm: {_fmt(r0.get('Tăng/giảm (kWh)',0))} kWh ({float(r0.get('Tăng/giảm (%)',0)):.2f}%)")
        if pcth_detail is not None and not pcth_detail.empty:
            cols=[c for c in ["Dòng dự báo","Loại","Điện nền (kWh)","Dự báo (kWh)","Tăng so nền (%)","Căn cứ"] if c in pcth_detail.columns]
            t=doc.add_table(rows=1, cols=len(cols))
            for i,c in enumerate(cols): t.rows[0].cells[i].text=c
            for _,row in pcth_detail.iterrows():
                cells=t.add_row().cells
                for i,c in enumerate(cols):
                    v=row[c]
                    cells[i].text=_fmt(v) if c in ["Điện nền (kWh)","Dự báo (kWh)"] else (f"{float(v):.2f}%" if c=="Tăng so nền (%)" and pd.notna(v) else str(v))
        if pcth_scenarios is not None and not pcth_scenarios.empty:
            doc.add_paragraph("Ba kịch bản là giả định thực hành, không phải xác suất/khoảng tin cậy.")
            cols=[c for c in ["Kịch bản","Tổng Điện lực (kWh)","Điều chỉnh g (điểm %)","Điều chỉnh nhiệt (°C)"] if c in pcth_scenarios.columns]
            t=doc.add_table(rows=1, cols=len(cols))
            for i,c in enumerate(cols): t.rows[0].cells[i].text=c
            for _,row in pcth_scenarios.iterrows():
                cells=t.add_row().cells
                for i,c in enumerate(cols): cells[i].text=_fmt(row[c]) if c=="Tổng Điện lực (kWh)" else str(row[c])
        if pcth_audit:
            doc.add_paragraph(f"Người lập: {pcth_audit.get('preparer','')} | Kiểm tra: {pcth_audit.get('checker','')} | Duyệt: {pcth_audit.get('approver','')}")
            doc.add_paragraph(f"Giải trình: {pcth_audit.get('reason','')}")
            doc.add_paragraph(f"Checklist trước trình duyệt: {'ĐÃ ĐỦ' if pcth_audit.get('ready') else 'CHƯA ĐỦ'}")

    doc.add_heading("4. Dự báo thống kê đối chiếu", level=1)
    table = doc.add_table(rows=1, cols=3)
    hdr = table.rows[0].cells
    hdr[0].text = "Tháng"; hdr[1].text = "Dự báo (kWh)"; hdr[2].text = "Mô hình"
    for _, row in forecasts.iterrows():
        cells = table.add_row().cells
        cells[0].text = pd.to_datetime(row['date']).strftime('%m/%Y')
        cells[1].text = _fmt(row['forecast'])
        cells[2].text = "Adaptive Ensemble"

    doc.add_heading("5. Kiểm định mô hình", level=1)
    if backtest is None or backtest.empty:
        doc.add_paragraph("Chưa đủ dữ liệu để kiểm định.")
    else:
        cols = [c for c in ["model","MAPE_1step","MAPE_2step","MAE_1step","RMSE_1step"] if c in backtest.columns]
        t = doc.add_table(rows=1, cols=len(cols))
        for i,c in enumerate(cols): t.rows[0].cells[i].text = c
        for _, row in backtest.iterrows():
            cells = t.add_row().cells
            for i,c in enumerate(cols):
                v=row[c]; cells[i].text = f"{v:.2f}" if isinstance(v,(float,int)) and c!="model" else str(v)

    if top100 is not None and not top100.empty:
        doc.add_heading("6. Top 10 khách hàng ảnh hưởng", level=1)
        use = top100.head(10)
        cols = [c for c in ["rank","customer_id","customer_name","latest_kwh","delta_kwh","share_pct","influence_score"] if c in use.columns]
        t = doc.add_table(rows=1, cols=len(cols))
        for i,c in enumerate(cols): t.rows[0].cells[i].text=c
        for _, row in use.iterrows():
            cells=t.add_row().cells
            for i,c in enumerate(cols): cells[i].text=_fmt(row[c]) if c in ["latest_kwh","delta_kwh"] else str(row[c])

    doc.add_paragraph("\nBáo cáo được tạo tự động bởi EVN Forecast 1.3 PCTH Web.")
    doc.save(bio)
    return bio.getvalue()


def build_pdf(series, forecasts, backtest, top100=None, state=None, outage_monthly=None, outage_details=None, pcth_detail=None, pcth_summary=None, pcth_scenarios=None, pcth_audit=None, title="EVN Forecast 1.3 PCTH"):
    bio = io.BytesIO()
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("t", parent=styles["Title"], alignment=TA_CENTER, fontSize=16, leading=20)
    doc = SimpleDocTemplate(bio, pagesize=A4, rightMargin=14*mm, leftMargin=14*mm, topMargin=14*mm, bottomMargin=14*mm)
    story = [Paragraph(title, title_style), Paragraph("Dien luc Thuong Xuan", styles["Heading2"]), Spacer(1,6)]
    latest = series.iloc[-1]
    story += [Paragraph(f"Du lieu moi nhat: {pd.to_datetime(latest['date']).strftime('%m/%Y')} - {_fmt(latest['actual'])} kWh", styles["BodyText"]), Spacer(1,6)]
    if "normalized_actual" in latest.index:
        story += [Paragraph(f"Chuan hoa mat dien: {_fmt(latest['normalized_actual'])} kWh; dien nang uoc khong thuc hien: {_fmt(latest.get('outage_lost_kwh',0))} kWh", styles["BodyText"]), Spacer(1,6)]
    if state:
        story += [Paragraph(f"Model State: EVNF_{state.get('latest_month') or 'NA'}_v1.2", styles["BodyText"]), Spacer(1,6)]

    if outage_monthly is not None and not outage_monthly.empty:
        story.append(Paragraph("Phan tich mat dien", styles["Heading2"]))
        rows=[["Thang","Dien nang uoc khong thuc hien (kWh)"]]
        for _,r in outage_monthly.iterrows(): rows.append([pd.to_datetime(r['date']).strftime('%m/%Y'),_fmt(r['lost_kwh'])])
        tbl=Table(rows,colWidths=[45*mm,80*mm]); tbl.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.lightgrey),('GRID',(0,0),(-1,-1),0.4,colors.grey)]))
        story += [tbl, Spacer(1,8)]


    if pcth_summary is not None and not pcth_summary.empty:
        story.append(Paragraph("PCTH Forecast Mode", styles["Heading2"]))
        r0=pcth_summary.iloc[0]
        story += [Paragraph(f"Tong nen: {_fmt(r0.get('Tổng nền (kWh)',0))} kWh; Du bao PCTH: {_fmt(r0.get('Tổng dự báo PCTH (kWh)',0))} kWh; Tang/giam: {float(r0.get('Tăng/giảm (%)',0)):.2f}%", styles["BodyText"]), Spacer(1,6)]
        if pcth_scenarios is not None and not pcth_scenarios.empty:
            rows=[["Kich ban","Tong Dien luc (kWh)"]]
            for _,r in pcth_scenarios.iterrows(): rows.append([str(r.get('Kịch bản','')),_fmt(r.get('Tổng Điện lực (kWh)',0))])
            tbl=Table(rows,colWidths=[45*mm,65*mm]); tbl.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.lightgrey),('GRID',(0,0),(-1,-1),0.4,colors.grey)]))
            story += [tbl, Spacer(1,8)]

    story.append(Paragraph("Du bao thong ke doi chieu", styles["Heading2"]))
    rows=[["Thang","Du bao (kWh)"]]
    for _,r in forecasts.iterrows(): rows.append([pd.to_datetime(r['date']).strftime('%m/%Y'),_fmt(r['forecast'])])
    tbl=Table(rows,colWidths=[45*mm,55*mm]); tbl.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.lightgrey),('GRID',(0,0),(-1,-1),0.4,colors.grey),('ALIGN',(1,1),(-1,-1),'RIGHT')]))
    story += [tbl, Spacer(1,10)]

    story.append(Paragraph("Kiem dinh mo hinh", styles["Heading2"]))
    if backtest is not None and not backtest.empty:
        rows=[["Model","MAPE 1","MAPE 2"]]
        for _,r in backtest.head(8).iterrows(): rows.append([str(r['model']),f"{r.get('MAPE_1step',0):.2f}%",f"{r.get('MAPE_2step',0):.2f}%"])
        tbl=Table(rows,colWidths=[60*mm,35*mm,35*mm]); tbl.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.lightgrey),('GRID',(0,0),(-1,-1),0.4,colors.grey)]))
        story += [tbl, Spacer(1,10)]

    if top100 is not None and not top100.empty:
        story.append(Paragraph("Top 10 khach hang anh huong", styles["Heading2"]))
        rows=[["STT","Khach hang","Moi nhat","Delta"]]
        for _,r in top100.head(10).iterrows(): rows.append([str(r.get('rank','')),str(r.get('customer_name',''))[:34],_fmt(r.get('latest_kwh',0)),_fmt(r.get('delta_kwh',0))])
        tbl=Table(rows,colWidths=[12*mm,85*mm,35*mm,35*mm]); tbl.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.lightgrey),('GRID',(0,0),(-1,-1),0.35,colors.grey),('FONTSIZE',(0,0),(-1,-1),7)]))
        story.append(tbl)
    story += [Spacer(1,10), Paragraph("Bao cao duoc tao tu dong boi EVN Forecast 1.3 PCTH Web.", styles["BodyText"])]
    doc.build(story)
    return bio.getvalue()
